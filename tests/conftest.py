import io
import os

# Must be set before the app is imported, so tests never load the real model at startup
os.environ["PRELOAD_MODEL"] = "false"
os.environ["GEMINI_API_KEY"] = ""

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.main import app
from core.config import get_settings
from core.scoring import ResumeAnalyzer, get_analyzer

VOCABULARY = ["python", "sql", "docker", "kubernetes", "dashboard", "marketing", "sales"]


class FakeEncoder:
    """Deterministic stand-in for SBERT: a normalised bag-of-words over a tiny vocabulary."""

    def encode(self, texts: list[str]) -> np.ndarray:
        vectors = np.array(
            [[text.lower().count(word) for word in VOCABULARY] for text in texts], dtype=float
        )
        norms = np.linalg.norm(vectors, axis=1, keepdims=True)
        return vectors / np.where(norms == 0, 1, norms)


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    if table:
        doc_table = doc.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(doc_table.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_pdf(text: str) -> bytes:
    """Builds a minimal single-page PDF containing one line of text."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    pdf = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (number, body)

    xref_offset = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        pdf += b"%010d 00000 n \n" % offset
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_offset,
    )
    return pdf


@pytest.fixture
def analyzer() -> ResumeAnalyzer:
    return ResumeAnalyzer(FakeEncoder())


@pytest.fixture
def client(analyzer):
    get_settings.cache_clear()
    app.dependency_overrides[get_analyzer] = lambda: analyzer
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    get_settings.cache_clear()
